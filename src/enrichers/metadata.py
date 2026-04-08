"""
File metadata extractor — universal first pass for all file types.
Extracts all available attributes using local tools only:
  - ffprobe (video/audio)
  - ExifRead / Pillow (images)
  - PyMuPDF (PDF)
  - python-docx / openpyxl (Office docs)
  - pefile / lief (executables)
  - cryptography (certificates)
  - py7zr / zipfile (archives)
  - Built-in (source code, data files)
"""
from __future__ import annotations

import json
import os
import subprocess
import zipfile
from pathlib import Path
from typing import Any


# ── File category detection ────────────────────────────────────────────────────

VIDEO_EXT      = {".mkv",".mp4",".mov",".avi",".ts",".m2ts",".wmv",".flv",".webm",".m4v"}
AUDIO_EXT      = {".flac",".mp3",".aac",".wav",".ogg",".m4a",".wma",".dsf",".dff",".opus"}
IMAGE_EXT      = {".jpg",".jpeg",".png",".heic",".tiff",".tif",".webp",".gif",".bmp",".raw",".cr2",".nef",".arw"}
DOC_EXT        = {".pdf",".docx",".doc",".xlsx",".xls",".pptx",".ppt",".odt",".ods",".odp",".txt",".md",".rtf"}
EXEC_EXT       = {".exe",".dll",".msi",".sys",".scr",".com"}
ARCHIVE_EXT    = {".zip",".rar",".7z",".tar",".gz",".bz2",".xz",".iso",".dmg"}
CERT_EXT       = {".pem",".crt",".cer",".p12",".pfx",".key",".der"}
CODE_EXT       = {".py",".js",".ts",".go",".rs",".cpp",".c",".h",".java",".cs",".rb",".php",".swift",".kt"}
WEB_EXT        = {".html",".htm",".json",".xml",".csv",".yaml",".yml",".sql",".graphql"}
EMAIL_EXT      = {".eml",".msg",".mbox"}
CAL_EXT        = {".ics",".vcf"}
MODEL_3D_EXT   = {".stl",".obj",".fbx",".blend",".step",".stp",".dwg",".dxf"}


def categorize(ext: str) -> str:
    ext = ext.lower()
    if ext in VIDEO_EXT:    return "video"
    if ext in AUDIO_EXT:    return "audio"
    if ext in IMAGE_EXT:    return "image"
    if ext in DOC_EXT:      return "document"
    if ext in EXEC_EXT:     return "executable"
    if ext in ARCHIVE_EXT:  return "archive"
    if ext in CERT_EXT:     return "certificate"
    if ext in CODE_EXT:     return "code"
    if ext in WEB_EXT:      return "web_data"
    if ext in EMAIL_EXT:    return "email"
    if ext in CAL_EXT:      return "calendar"
    if ext in MODEL_3D_EXT: return "3d_model"
    return "other"


# ── Extractors ────────────────────────────────────────────────────────────────

def extract_all(file_path: str) -> dict[str, Any]:
    """
    Run all applicable extractors for a file.
    Returns a flat dict of all extracted attributes.
    """
    ext = Path(file_path).suffix.lower()
    category = categorize(ext)
    props: dict[str, Any] = {"file_category": category}

    try:
        if category == "video":
            props.update(_extract_video(file_path))
        elif category == "audio":
            props.update(_extract_audio(file_path))
        elif category == "image":
            props.update(_extract_image(file_path))
        elif category == "document":
            props.update(_extract_document(file_path, ext))
        elif category == "executable":
            props.update(_extract_executable(file_path))
        elif category == "archive":
            props.update(_extract_archive(file_path, ext))
        elif category == "certificate":
            props.update(_extract_certificate(file_path))
        elif category == "code":
            props.update(_extract_code(file_path))
        elif category == "web_data":
            props.update(_extract_web_data(file_path, ext))
    except Exception as e:
        props["extraction_error"] = str(e)

    return props


def _ffprobe(file_path: str) -> dict:
    result = subprocess.run(
        ["ffprobe", "-v", "quiet", "-print_format", "json",
         "-show_format", "-show_streams", file_path],
        capture_output=True, text=True, timeout=30
    )
    return json.loads(result.stdout) if result.returncode == 0 else {}


def _extract_video(path: str) -> dict[str, Any]:
    data = _ffprobe(path)
    props: dict[str, Any] = {}
    fmt = data.get("format", {})
    props["duration_secs"] = float(fmt.get("duration", 0)) or None
    props["overall_bitrate"] = int(fmt.get("bit_rate", 0)) or None
    props["container_format"] = fmt.get("format_long_name")

    for stream in data.get("streams", []):
        ctype = stream.get("codec_type")
        if ctype == "video" and "video_codec" not in props:
            props["video_codec"] = stream.get("codec_name")
            props["width"] = stream.get("width")
            props["height"] = stream.get("height")
            fr = stream.get("avg_frame_rate", "0/1").split("/")
            props["fps"] = round(int(fr[0]) / max(int(fr[1]), 1), 3) if len(fr) == 2 else None
            props["bit_depth"] = stream.get("bits_per_raw_sample")
            props["color_space"] = stream.get("color_space")
            ct = stream.get("color_transfer", "")
            if "smpte2084" in ct or "pq" in ct:
                props["hdr_format"] = "HDR10"
            elif "arib-std-b67" in ct or "hlg" in ct:
                props["hdr_format"] = "HLG"
        elif ctype == "audio" and "audio_codec" not in props:
            props["audio_codec"] = stream.get("codec_name")
            props["audio_channels"] = stream.get("channels")
            props["sample_rate"] = int(stream.get("sample_rate", 0)) or None
        elif ctype == "subtitle":
            langs = props.get("subtitle_languages", "")
            lang = stream.get("tags", {}).get("language", "")
            if lang:
                props["subtitle_languages"] = f"{langs},{lang}".strip(",")

    return {k: v for k, v in props.items() if v is not None}


def _extract_audio(path: str) -> dict[str, Any]:
    data = _ffprobe(path)
    props: dict[str, Any] = {}
    fmt = data.get("format", {})
    tags = fmt.get("tags", {})
    props["duration_secs"] = float(fmt.get("duration", 0)) or None
    props["overall_bitrate"] = int(fmt.get("bit_rate", 0)) or None

    for stream in data.get("streams", []):
        if stream.get("codec_type") == "audio":
            props["audio_codec"] = stream.get("codec_name")
            props["audio_channels"] = stream.get("channels")
            props["sample_rate"] = int(stream.get("sample_rate", 0)) or None
            props["bit_depth"] = stream.get("bits_per_raw_sample")
            break

    # ID3/Vorbis tags
    def t(key: str) -> str | None:
        return tags.get(key) or tags.get(key.upper()) or None

    props["artist"]       = t("artist")
    props["album"]        = t("album")
    props["album_artist"] = t("album_artist")
    props["title"]        = t("title")
    props["genre"]        = t("genre")
    props["track_number"] = _safe_int(t("track"))
    props["bpm"]          = _safe_float(t("bpm") or t("TBPM"))
    props["musicbrainz_id"] = t("musicbrainz_trackid")
    try:
        props["year"] = int(str(t("date") or "")[:4]) if t("date") else None
    except Exception:
        pass

    return {k: v for k, v in props.items() if v is not None}


def _extract_image(path: str) -> dict[str, Any]:
    props: dict[str, Any] = {}
    try:
        import exifread
        with open(path, "rb") as f:
            tags = exifread.process_file(f, details=False)

        def exif(key: str):
            v = tags.get(key)
            return str(v) if v else None

        props["camera_make"]  = exif("Image Make")
        props["camera_model"] = exif("Image Model")
        props["lens"]         = exif("EXIF LensModel")
        props["focal_length"] = _safe_float(exif("EXIF FocalLength"))
        props["aperture"]     = _safe_float(exif("EXIF FNumber"))
        props["shutter_speed"]= exif("EXIF ExposureTime")
        props["iso"]          = _safe_int(exif("EXIF ISOSpeedRatings"))

        # GPS
        lat = _parse_gps(tags.get("GPS GPSLatitude"), tags.get("GPS GPSLatitudeRef"))
        lon = _parse_gps(tags.get("GPS GPSLongitude"), tags.get("GPS GPSLongitudeRef"))
        if lat: props["gps_latitude"]  = lat
        if lon: props["gps_longitude"] = lon

        # Datetime
        dt_str = exif("EXIF DateTimeOriginal") or exif("Image DateTime")
        if dt_str:
            try:
                from datetime import datetime
                props["datetime_original"] = datetime.strptime(
                    dt_str, "%Y:%m:%d %H:%M:%S"
                ).timestamp()
            except Exception:
                pass
    except Exception:
        pass

    # Dimensions via Pillow
    try:
        from PIL import Image as PILImage
        with PILImage.open(path) as img:
            props["width"], props["height"] = img.size
            props["color_profile"] = img.mode
            props["bit_depth"] = {"1": 1, "L": 8, "RGB": 24, "RGBA": 32}.get(img.mode)
    except Exception:
        pass

    return {k: v for k, v in props.items() if v is not None}


def _extract_document(path: str, ext: str) -> dict[str, Any]:
    props: dict[str, Any] = {}

    if ext == ".pdf":
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(path)
            meta = doc.metadata
            props["page_count"]   = doc.page_count
            props["author"]       = meta.get("author")
            props["title"]        = meta.get("title")
            props["is_encrypted"] = doc.is_encrypted
            # Word count estimate
            text = "".join(page.get_text() for page in doc)
            props["word_count"] = len(text.split())
            doc.close()
        except Exception:
            pass

    elif ext in {".docx", ".doc"}:
        try:
            from docx import Document
            doc = Document(path)
            core = doc.core_properties
            props["author"]     = core.author
            props["title"]      = core.title
            props["word_count"] = sum(len(p.text.split()) for p in doc.paragraphs)
            props["page_count"] = None  # python-docx doesn't expose page count
        except Exception:
            pass

    elif ext in {".xlsx", ".xls"}:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            props["page_count"] = len(wb.sheetnames)  # sheets
            props["author"]     = wb.properties.creator
        except Exception:
            pass

    # Secret detection placeholder
    try:
        text_path = path if ext in {".txt", ".md", ".csv"} else None
        if text_path:
            with open(text_path, "r", errors="replace") as f:
                content = f.read(50000)
            if any(kw in content.lower() for kw in
                   ["api_key", "secret", "password", "token", "private_key"]):
                props["contains_secrets"] = True
    except Exception:
        pass

    return {k: v for k, v in props.items() if v is not None}


def _extract_executable(path: str) -> dict[str, Any]:
    """Cross-platform binary analysis via lief (PE/ELF/Mach-O)."""
    try:
        from src.extractors.binary import analyze_binary
        return analyze_binary(path)
    except Exception as e:
        return {"extraction_error": str(e)}


def _extract_archive(path: str, ext: str) -> dict[str, Any]:
    props: dict[str, Any] = {}
    try:
        if ext == ".zip":
            with zipfile.ZipFile(path) as z:
                names = z.namelist()
                props["file_count_in_archive"] = len(names)
                total_compressed = sum(i.compress_size for i in z.infolist())
                total_original = sum(i.file_size for i in z.infolist())
                props["compression_ratio"] = round(
                    total_compressed / max(total_original, 1), 3
                )
                props["contains_executables"] = any(
                    n.lower().endswith((".exe", ".dll", ".bat", ".ps1"))
                    for n in names
                )
        elif ext == ".7z":
            try:
                import py7zr
                with py7zr.SevenZipFile(path) as z:
                    props["file_count_in_archive"] = len(z.list())
            except Exception:
                pass
    except Exception:
        pass
    return {k: v for k, v in props.items() if v is not None}


def _extract_certificate(path: str) -> dict[str, Any]:
    props: dict[str, Any] = {}
    try:
        from cryptography import x509
        from cryptography.hazmat.primitives import serialization
        from datetime import datetime, timezone

        with open(path, "rb") as f:
            data = f.read()

        try:
            cert = x509.load_pem_x509_certificate(data)
        except Exception:
            try:
                cert = x509.load_der_x509_certificate(data)
            except Exception:
                return props

        now = datetime.now(timezone.utc)
        props["cert_subject"]     = cert.subject.rfc4514_string()
        props["cert_issuer"]      = cert.issuer.rfc4514_string()
        props["cert_valid_from"]  = cert.not_valid_before_utc.timestamp()
        props["cert_valid_to"]    = cert.not_valid_after_utc.timestamp()
        props["cert_is_expired"]  = cert.not_valid_after_utc < now
        props["days_until_expiry"]= (cert.not_valid_after_utc - now).days
        props["cert_fingerprint"] = cert.fingerprint(
            __import__("cryptography.hazmat.primitives.hashes", fromlist=["SHA256"]).SHA256()
        ).hex()
        props["is_self_signed"]   = cert.issuer == cert.subject
    except Exception:
        pass
    return {k: v for k, v in props.items() if v is not None}


def _extract_code(path: str) -> dict[str, Any]:
    props: dict[str, Any] = {}
    try:
        with open(path, "r", errors="replace") as f:
            lines = f.readlines()
        props["line_count"] = len(lines)
        content = "".join(lines)
        # Secret heuristics
        secret_patterns = [
            "api_key", "apikey", "secret_key", "access_token",
            "private_key", "password", "passwd", "auth_token",
            "bearer ", "-----begin rsa"
        ]
        found = [p for p in secret_patterns if p in content.lower()]
        if found:
            props["contains_secrets"] = True
            props["secret_types"] = ",".join(found)

        # Language-specific stats
        ext = Path(path).suffix.lower()
        props["code_language"] = {
            ".py": "Python", ".js": "JavaScript", ".ts": "TypeScript",
            ".go": "Go", ".rs": "Rust", ".cpp": "C++", ".c": "C",
            ".java": "Java", ".cs": "C#", ".rb": "Ruby",
        }.get(ext, ext.lstrip(".").upper())

        # Function count heuristic
        markers = {
            ".py": "def ", ".js": "function ", ".ts": "function ",
            ".go": "func ", ".rs": "fn ", ".java": "void |public |private |protected ",
        }
        marker = markers.get(ext)
        if marker:
            props["function_count"] = sum(1 for l in lines if marker.split("|")[0] in l)
    except Exception:
        pass
    return {k: v for k, v in props.items() if v is not None}


def _extract_web_data(path: str, ext: str) -> dict[str, Any]:
    props: dict[str, Any] = {}
    try:
        with open(path, "r", errors="replace") as f:
            content = f.read(100000)
        props["line_count"] = content.count("\n")

        if ext == ".csv":
            lines = content.split("\n")
            if lines:
                props["page_count"] = len(lines)  # rows
                props["word_count"] = len(lines[0].split(","))  # columns

        # PII detection heuristic
        import re
        pii_found = []
        if re.search(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b", content):
            pii_found.append("email")
        if re.search(r"\b\d{3}[-.]?\d{3}[-.]?\d{4}\b", content):
            pii_found.append("phone")
        if re.search(r"\b\d{3}-\d{2}-\d{4}\b", content):
            pii_found.append("ssn")
        if pii_found:
            props["pii_detected"] = True
            props["pii_types"] = ",".join(pii_found)
            props["sensitivity_level"] = "high"
    except Exception:
        pass
    return {k: v for k, v in props.items() if v is not None}


# ── Helpers ────────────────────────────────────────────────────────────────────

def _safe_int(v) -> int | None:
    try: return int(str(v).split("/")[0])
    except: return None

def _safe_float(v) -> float | None:
    try:
        s = str(v)
        if "/" in s:
            parts = s.split("/")
            return round(float(parts[0]) / float(parts[1]), 4)
        return float(s)
    except: return None

def _parse_gps(coord, ref) -> float | None:
    try:
        vals = [float(str(v)) for v in coord.values]
        decimal = vals[0] + vals[1]/60 + vals[2]/3600
        if str(ref) in ("S", "W"):
            decimal = -decimal
        return round(decimal, 6)
    except: return None
